#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Generador Masivo de Reportes de Pago a Proveedores (SAP)
Procesa carpetas por departamento, extrae datos de Consolidados y fechas de Órdenes de Compra (PDF),
y genera el Excel consolidado 'Masivo Reporte.xlsx' con formato oficial columnas A-U.
"""

import os
import sys
import re
import glob

# Asegurar compatibilidad UTF-8 en consola Windows
if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')
if hasattr(sys.stderr, 'reconfigure'):
    sys.stderr.reconfigure(encoding='utf-8')

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

# ── Estilos para Excel ──────────────────────────────────────────
thin = Side(style='thin', color='BFBFBF')
thin_b = Border(left=thin, right=thin, top=thin, bottom=thin)

fill_hdr_green = PatternFill('solid', fgColor='375623')
fill_hdr_teal  = PatternFill('solid', fgColor='00B0F0')
fill_hdr_olive = PatternFill('solid', fgColor='70AD47')

font_hdr_white = Font(name='Calibri', bold=True, size=10, color='FFFFFF')
font_hdr_black = Font(name='Calibri', bold=True, size=10, color='000000')

fill_white   = PatternFill('solid', fgColor='FFFFFF')
fill_alt     = PatternFill('solid', fgColor='F2F2F2')
fill_oc      = PatternFill('solid', fgColor='DEEAF1')
fill_importe = PatternFill('solid', fgColor='FFF2CC')

font_data  = Font(name='Calibri', size=10)
font_bold  = Font(name='Calibri', size=10, bold=True)
font_oc    = Font(name='Calibri', size=10, bold=True, color='1F4E79')

al_c  = Alignment(horizontal='center', vertical='center')
al_l  = Alignment(horizontal='left',   vertical='center')
al_r  = Alignment(horizontal='right',  vertical='center')
al_wc = Alignment(horizontal='center', vertical='center', wrap_text=True)

COL_DEFS = [
    (1,  'A', 'ID CABC(1)',               8,    fill_hdr_green, font_hdr_white, al_c, None),
    (2,  'B', 'Fecha Doc',               13,    fill_hdr_green, font_hdr_white, al_c, None),
    (3,  'C', 'Texto Cab Doc',           20,    fill_hdr_olive, font_hdr_white, al_l, None),
    (4,  'D', 'REFERENCIA',              22,    fill_hdr_green, font_hdr_white, al_c, None),
    (5,  'E', 'ID POSC(2)Acreedor',      12,    fill_hdr_olive, font_hdr_white, al_c, None),
    (6,  'F', 'RUC',                     14,    fill_hdr_olive, font_hdr_white, al_c, None),
    (7,  'G', 'Cond Pago',               10,    fill_hdr_olive, font_hdr_white, al_c, None),
    (8,  'H', 'Ind Impto',               10,    fill_hdr_olive, font_hdr_white, al_c, None),
    (9,  'I', 'Tipo Detr',               10,    fill_hdr_olive, font_hdr_white, al_c, None),
    (10, 'J', 'Ind Detra',               10,    fill_hdr_olive, font_hdr_white, al_c, None),
    (11, 'K', 'Asignación',              14,    fill_hdr_olive, font_hdr_white, al_c, None),
    (12, 'L', 'Texto Posición',          42,    fill_hdr_teal,  font_hdr_black, al_l, None),
    (13, 'M', 'ID POSC(2) Cuenta GASTO', 10,    fill_hdr_olive, font_hdr_white, al_c, None),
    (14, 'N', 'Cuenta Gasto',            13,    fill_hdr_olive, font_hdr_white, al_c, None),
    (15, 'O', 'DETALLE',                 20,    fill_hdr_olive, font_hdr_white, al_l, None),
    (16, 'P', 'Centro Costo',            14,    fill_hdr_olive, font_hdr_white, al_c, None),
    (17, 'Q', 'Importe MD',              14,    fill_hdr_green, font_hdr_white, al_r, '"S/ "#,##0.00'),
    (18, 'R', 'Ind Impto',               10,    fill_hdr_olive, font_hdr_white, al_c, None),
    (19, 'S', 'Asignación',              14,    fill_hdr_olive, font_hdr_white, al_c, None),
    (20, 'T', 'Texto Posición',          42,    fill_hdr_teal,  font_hdr_black, al_l, None),
    (21, 'U', 'N° OC',                   14,    fill_hdr_teal,  font_hdr_black, al_c, None),
]

MESES_PE = {
    'enero':'01', 'febrero':'02', 'marzo':'03', 'abril':'04', 'mayo':'05', 'junio':'06',
    'julio':'07', 'agosto':'08', 'setiembre':'09', 'septiembre':'09', 'octubre':'10',
    'noviembre':'11', 'diciembre':'12'
}

def escanear_documentos_pdf(dir_path):
    """
    Escanea todos los PDFs de la carpeta detectando:
    - Formato Imagen 1: Factura Electrónica (Importe Total, IGV, Subtotal)
    - Formato Imagen 2: Recibo por Honorarios Electrónico (Total Honorarios / Neto)
    - Órdenes de Compra (OC): Fechas de emisión
    """
    doc_map = {}
    oc_dates = {}

    pdf_files = [f for f in os.listdir(dir_path) if f.upper().endswith('.PDF')]
    for fname in pdf_files:
        fpath = os.path.join(dir_path, fname)
        txt = ""
        # 1. Intentar extracción con pypdf
        try:
            import pypdf
            reader = pypdf.PdfReader(fpath)
            txt = ' '.join(page.extract_text() or '' for page in reader.pages)
        except Exception:
            pass

        # 2. Fallback decodificación binaria / hex tokens si txt es corto
        if len(txt.strip()) < 15:
            try:
                with open(fpath, 'rb') as fp:
                    raw = fp.read()
                tokens = re.findall(rb'<([0-9A-Fa-f]{4,})>', raw)
                decoded = []
                for t in tokens:
                    try:
                        decoded.append(bytes.fromhex(t.decode()).decode('latin1'))
                    except Exception:
                        pass
                txt += ' ' + ' '.join(decoded)
            except Exception:
                pass

        clean_txt = re.sub(r'[\x00-\x08\x0B\x0C\x0E-\x1F\ufffd\xa0]', ' ', txt)
        clean_txt = re.sub(r'\s+', ' ', clean_txt)
        txt_upper = clean_txt.upper()

        fn_upper = fname.upper()

        # Descartar archivos que son HES, Solped u Orden de Compra antes de clasificar comprobantes
        is_hes_oc_solped = bool(re.search(r'\b(HES|SOLPED|PEDIDO|OC)\b', fn_upper)) or ('ORDEN DE' in txt_upper)

        # Detección Formato Imagen 1: Factura Electrónica
        is_factura = (not is_hes_oc_solped) and (
            fn_upper.startswith('FACT') or
            ('FACTURA ELECTRONICA' in txt_upper) or
            ('FACTURA' in txt_upper and ('IMPORTE TOTAL' in txt_upper or 'SUB TOTAL' in txt_upper))
        )

        # Detección Formato Imagen 2: Recibo por Honorarios
        is_rxh = (not is_hes_oc_solped) and (
            fn_upper.startswith('RXH') or
            ('RECIBO POR HONORARIOS ELECTRONICO' in txt_upper) or
            ('RECIBO POR HONORARIOS' in txt_upper)
        )

        # Detección Orden de Compra
        is_oc = bool(re.search(r'\bOC\b', fn_upper)) or ('ORDEN DE' in txt_upper)

        # Extraer clave de documento (ej: E001-1388 -> E0011388, E001 134 -> E001134)
        m_num = re.search(r'([EFEB]\d{3})[-\s]*(\d+)', fname + ' ' + clean_txt, re.IGNORECASE)
        num_key = (m_num.group(1) + m_num.group(2)).upper() if m_num else None

        # Extraer RUC del PDF
        m_ruc = re.search(r'R\.?U\.?C\.?\s*[:\-]?\s*(\d{11})', clean_txt, re.IGNORECASE)
        pdf_ruc = m_ruc.group(1) if m_ruc else None

        # Extraer Importe
        total_monto = None
        if is_factura:
            # Factura: elegir el Importe Total (con IGV)
            m_tot = re.search(r'Importe\s+Total\s*:\s*(?:S/?\s*)?([0-9,]+\.[0-9]{2})', clean_txt, re.I)
            if m_tot:
                total_monto = float(m_tot.group(1).replace(',', ''))
        elif is_rxh:
            # Recibo por Honorario: Total Neto Recibido o Total por honorarios
            m_tot = re.search(r'Total\s+Neto\s+Recibido\s*:\s*(?:SOLES\s*)?([0-9,]+\.[0-9]{2})', clean_txt, re.I) or \
                    re.search(r'Total\s+por\s+honorarios\s*:\s*([0-9,]+\.[0-9]{2})', clean_txt, re.I)
            if m_tot:
                total_monto = float(m_tot.group(1).replace(',', ''))
            else:
                m_num_after = re.search(r'\b202\d\s+([0-9,]+\.[0-9]{2})', clean_txt)
                if m_num_after:
                    total_monto = float(m_num_after.group(1).replace(',', ''))

        # Extraer Fecha de Emisión
        fecha_emision = None
        m_date = re.search(r'Fecha\s+de\s+[eE]misi[oó]n\s*:\s*(\d{2}[./-]\d{2}[./-]\d{4})', clean_txt, re.I)
        if m_date:
            fecha_emision = m_date.group(1).replace('/', '.').replace('-', '.')
        else:
            m_date2 = re.search(r'(\d{1,2})\s+(?:de\s+)?([A-Za-z]+)\s+(?:del?\s+)?(\d{4})', clean_txt, re.I)
            if m_date2:
                m_mes = m_date2.group(2).lower()
                if m_mes in MESES_PE:
                    fecha_emision = f"{int(m_date2.group(1)):02d}.{MESES_PE[m_mes]}.{m_date2.group(3)}"

        # Guardar en mapas
        doc_info = {
            'tipo': 'FACTURA' if is_factura else 'RXH' if is_rxh else 'OC' if is_oc else 'OTRO',
            'num_key': num_key,
            'ruc': pdf_ruc,
            'total': total_monto,
            'fecha': fecha_emision,
            'fname': fname
        }

        if is_factura or is_rxh:
            if num_key:
                doc_map[num_key] = doc_info
            doc_map[fname.upper()] = doc_info

        # Fechas de OC
        if is_oc and fecha_emision:
            m_oc = re.search(r'\bOC[\s_\-]*(\d+)', fname, re.IGNORECASE)
            m_rxh = re.search(r'\b(?:RXH|FACTURA)[\s_\-]*([A-Za-z0-9\-]+)', fname, re.IGNORECASE)
            if m_oc:
                oc_dates['OC_' + m_oc.group(1).strip()] = fecha_emision
            if m_rxh:
                oc_dates['RXH_' + m_rxh.group(1).strip().upper()] = fecha_emision
            if '__DEFAULT__' not in oc_dates:
                oc_dates['__DEFAULT__'] = fecha_emision

    return doc_map, oc_dates

def normalizar_texto_pos(motivo, semana):
    """Estandariza abreviaturas según la regla de negocio."""
    t = str(motivo or '').strip()
    repl = [
        (r'\bSERV\.\s*', 'SERVICIO DE '),
        (r'\bTRANSP\.\s*', 'TRANSPORTE '),
        (r'\bMTTO\b', 'MANTENIMIENTO'),
        (r'\bAYUDANTE\s+REP\.\s*', 'AYUDANTE DE REPARTO '),
        (r'\bREP\.\s*', 'REPARTO '),
    ]
    for pattern, r in repl:
        t = re.sub(pattern, r, t, flags=re.IGNORECASE)
    t = re.sub(r'\s{2,}', ' ', t)
    t = re.sub(r'\s*SEMANA\s*\d+\s*$', '', t, flags=re.IGNORECASE).strip().upper()
    return f"{t} SEMANA {semana}"

def formatear_referencia(doc_num, tipo='RXH'):
    """
    Formatea la referencia para SAP según tipo de comprobante:
    - Factura (Formato Imagen 1): 01-0{SERIE}-{NUMERO_7_DIGITOS}
    - Recibo por Honorarios (Formato Imagen 2): 02-0{SERIE}-{NUMERO_7_DIGITOS}
    """
    s = str(doc_num or '').strip()
    prefix = '01-0' if str(tipo).upper() in ('FACTURA', 'FACT', 'FAC', '01') else '02-0'
    if s.startswith('01-0') or s.startswith('02-0'):
        return s
    s = re.sub(r'^[^\w]+', '', s)
    m = re.match(r'^([A-Za-z0-9]{3,4})[\s\-]+(\d+)$', s)
    if m:
        return f"{prefix}{m.group(1).upper()}-{m.group(2).zfill(7)}"
    if '-' in s:
        parts = s.split('-', 1)
        return f"{prefix}{parts[0].strip().upper()}-{parts[1].strip().zfill(7)}"
    return prefix + s

def procesar_departamento(dir_path, wb_out):
    """Procesa una subcarpeta de departamento y agrega una hoja al workbook."""
    folder_name = os.path.basename(dir_path)
    dept_clean = re.sub(r'^\d+[\s_\-]*', '', folder_name)
    dept_clean = re.sub(r'[\s_\-]*(?:SEM(?:ANA)?[\s_\-]*\d+).*$', '', dept_clean, flags=re.IGNORECASE).strip().upper()
    print(f"\n📂 Procesando carpeta: {folder_name} (Depto: {dept_clean})")

    # Buscar consolidado Excel
    excel_files = [f for f in os.listdir(dir_path) if f.lower().endswith(('.xlsx', '.xls')) and 'CONSOLIDADO' in f.upper()]
    if not excel_files:
        print("  ⚠️ No se encontró archivo CONSOLIDADO. Se omite la carpeta.")
        return False

    cons_path = os.path.join(dir_path, excel_files[0])
    print(f"  📄 Consolidado encontrado: {excel_files[0]}")

    # Extraer documentos PDF (Facturas, RH, OCs)
    doc_map, oc_dates = escanear_documentos_pdf(dir_path)
    print(f"  🔍 Documentos PDF detectados: {len(doc_map)} comprobantes, {len(oc_dates)} fechas OC")

    # Leer Consolidado
    wb_in = openpyxl.load_workbook(cons_path, data_only=True)
    sheet_name = next((s for s in wb_in.sheetnames if 'consolidado' in s.lower()), wb_in.sheetnames[0])
    ws_in = wb_in[sheet_name]

    # Detectar encabezados
    hdr_row = 1
    for r in range(1, min(6, ws_in.max_row + 1)):
        vals = [ws_in.cell(r, c).value for c in range(1, min(20, ws_in.max_column + 1))]
        if sum(1 for v in vals if v is not None) >= 5:
            hdr_row = r
            break

    headers = [str(ws_in.cell(hdr_row, c).value or '').upper().replace('°', '').replace(' ', '') for c in range(1, ws_in.max_column + 1)]
    def find_col(*terms):
        for t in terms:
            for idx, h in enumerate(headers):
                if t in h:
                    return idx + 1
        return -1

    c_id      = find_col('ID')
    c_fecha   = find_col('FECHA')
    c_sem     = find_col('SEM')
    c_prov    = find_col('PROVINCIA', 'PROV')
    c_motivo  = find_col('MOTIVO')
    c_ruc     = find_col('RUC')
    c_tipodoc = find_col('TIPODOC', 'TIPO')
    c_rxh     = find_col('RXH', 'NOC', 'N°')
    c_oc      = find_col('NOC', 'OC')

    # Filtrar filas válidas
    raw_rows = []
    for r in range(hdr_row + 1, ws_in.max_row + 1):
        prov = str(ws_in.cell(r, c_prov).value or '').upper().strip() if c_prov > 0 else ''
        if dept_clean and prov and (dept_clean in prov or prov in dept_clean):
            raw_rows.append(r)
        elif not dept_clean and any(ws_in.cell(r, c).value is not None for c in range(1, 15)):
            raw_rows.append(r)

    if not raw_rows:
        print("  ⚠️ No se encontraron filas para este departamento.")
        return False

    # Filtrar por semana más reciente
    semana = '?'
    if c_sem > 0:
        semanas = [ws_in.cell(r, c_sem).value for r in raw_rows if isinstance(ws_in.cell(r, c_sem).value, (int, float))]
        if semanas:
            semana = int(max(semanas))
            raw_rows = [r for r in raw_rows if ws_in.cell(r, c_sem).value == semana]

    print(f"  ✅ Registros filtrados: {len(raw_rows)} (Semana {semana})")

    # Crear hoja en wb_out
    title_dept = dept_clean.title()
    sheet_title = f"{title_dept} Semana {semana}"[:31]
    ws_out = wb_out.create_sheet(title=sheet_title)
    ws_out.row_dimensions[1].height = 32

    # Escribir encabezados
    for (col_idx, col_letter, header, width, hdr_fill, hdr_font, dat_al, num_fmt) in COL_DEFS:
        cell = ws_out.cell(row=1, column=col_idx, value=header)
        cell.fill = hdr_fill
        cell.font = hdr_font
        cell.alignment = al_wc
        cell.border = Border(left=Side(style='medium', color='FFFFFF'),
                             right=Side(style='medium', color='FFFFFF'),
                             top=thin, bottom=thin)
        ws_out.column_dimensions[col_letter].width = width

    # Escribir datos
    for seq, r in enumerate(raw_rows, start=1):
        row_out = seq + 1
        is_odd = (seq % 2 == 1)
        row_fill = fill_white if is_odd else fill_alt

        motivo   = ws_in.cell(r, c_motivo).value if c_motivo > 0 else ''
        rxh      = ws_in.cell(r, c_rxh).value if c_rxh > 0 else ''
        oc       = str(ws_in.cell(r, c_oc).value or '').strip() if c_oc > 0 else ''
        ruc      = str(ws_in.cell(r, c_ruc).value or '').strip() if c_ruc > 0 else ''
        tipodoc  = str(ws_in.cell(r, c_tipodoc).value or '').strip().upper() if c_tipodoc > 0 else ''
        fecha_orig = ws_in.cell(r, c_fecha).value if c_fecha > 0 else None

        tp = normalizar_texto_pos(motivo, semana)

        # Buscar coincidencia con PDF escaneado
        rxh_clean = re.sub(r'[\s\-]', '', str(rxh or '')).upper()
        matched_pdf = doc_map.get(rxh_clean)
        if not matched_pdf and rxh_clean:
            # Búsqueda difusa en el mapa de documentos
            for k, doc in doc_map.items():
                if k in rxh_clean or rxh_clean in k:
                    matched_pdf = doc
                    break

        # Determinar si es FACTURA (Formato Imagen 1) o RECIBO POR HONORARIOS (Formato Imagen 2)
        rxh_upper = str(rxh or '').strip().upper()
        if any(k in tipodoc for k in ('FACTURA', 'FACT', 'FAC', '01')):
            es_factura = True
        elif matched_pdf and matched_pdf.get('tipo') == 'FACTURA':
            es_factura = True
        elif rxh_upper.startswith('F') or ('FACT' in rxh_upper):
            es_factura = True
        elif ruc.startswith('20') and not any(k in tipodoc for k in ('RXH', 'RH', 'HONORARIO')):
            es_factura = True
        elif matched_pdf and matched_pdf.get('tipo') == 'RXH':
            es_factura = False
        elif any(k in tipodoc for k in ('RXH', 'RH', 'HONORARIO')):
            es_factura = False
        else:
            es_factura = False

        tipo_final = 'FACTURA' if es_factura else 'RXH'

        # Asignar Importe: EN CASO DE FACTURA ELEGIR EL IMPORTE TOTAL
        if es_factura:
            if matched_pdf and matched_pdf.get('total') is not None:
                val_num = matched_pdf['total']
            else:
                val_num = None
                print(f"⚠️ No se pudo leer el Importe Total de la factura {rxh}; VALORIZADO quedará vacío.")
            ind_impto_code = 'C1'  # 18% IGV en SAP
        else:
            if matched_pdf and matched_pdf.get('total') is not None:
                val_num = matched_pdf['total']
            else:
                val_num = None
                print(f"⚠️ No se pudo leer el importe neto del RH {rxh}; VALORIZADO quedará vacío.")
            ind_impto_code = 'C0'  # Inafecto / 0% IGV en SAP

        # Resolver fecha doc (Prioridad: PDF de Factura/RH -> OC PDF -> Consolidado)
        fdoc = None
        if matched_pdf and matched_pdf.get('fecha'):
            fdoc = matched_pdf['fecha']
        elif oc and ('OC_' + oc) in oc_dates:
            fdoc = oc_dates['OC_' + oc]
        elif rxh and ('RXH_' + str(rxh).strip().upper()) in oc_dates:
            fdoc = oc_dates['RXH_' + str(rxh).strip().upper()]
        elif '__DEFAULT__' in oc_dates:
            fdoc = oc_dates['__DEFAULT__']
        elif hasattr(fecha_orig, 'strftime'):
            fdoc = fecha_orig.strftime('%d.%m.%Y')
        else:
            fdoc = str(fecha_orig or '')

        # Referencia formateada: 01- para Factura, 02- para RH
        ref_formateada = formatear_referencia(rxh, tipo=tipo_final)

        cols_data = [
            (1,  seq,                                  al_c, None,            row_fill,    font_data),
            (2,  fdoc,                                 al_c, None,            row_fill,    font_data),
            (3,  None,                                 al_l, None,            row_fill,    font_data),
            (4,  ref_formateada,                       al_c, None,            row_fill,    font_data),
            (5,  2,                                    al_c, None,            row_fill,    font_data),
            (6,  ruc,                                  al_c, None,            row_fill,    font_data),
            (7,  'CP02',                               al_c, None,            row_fill,    font_data),
            (8,  ind_impto_code,                       al_c, None,            row_fill,    font_data),
            (9,  None,                                 al_c, None,            row_fill,    font_data),
            (10, None,                                 al_c, None,            row_fill,    font_data),
            (11, None,                                 al_c, None,            row_fill,    font_data),
            (12, tp,                                   al_l, None,            row_fill,    font_data),
            (13, 2,                                    al_c, None,            row_fill,    font_data),
            (14, None,                                 al_c, None,            row_fill,    font_data),
            (15, None,                                 al_l, None,            row_fill,    font_data),
            (16, None,                                 al_c, None,            row_fill,    font_data),
            (17, val_num,                              al_r, '"S/ "#,##0.00', fill_importe, font_bold),
            (18, ind_impto_code,                       al_c, None,            row_fill,    font_data),
            (19, None,                                 al_c, None,            row_fill,    font_data),
            (20, tp,                                   al_l, None,            row_fill,    font_data),
            (21, oc if oc else None,                   al_c, None,            fill_oc,     font_oc),
        ]

        for (c_idx, val, alignment, num_fmt, c_fill, c_font) in cols_data:
            cell = ws_out.cell(row=row_out, column=c_idx, value=val)
            cell.font = c_font
            cell.alignment = alignment
            cell.border = thin_b
            cell.fill = c_fill
            if num_fmt:
                cell.number_format = num_fmt

        ws_out.row_dimensions[row_out].height = 18

    # Fila de totales
    tot_row = len(raw_rows) + 2
    fill_tot = PatternFill('solid', fgColor='375623')
    font_tot = Font(name='Calibri', bold=True, size=10, color='FFFFFF')
    font_tot_q = Font(name='Calibri', bold=True, size=10, color='FFD700')

    for c in range(1, 22):
        cell = ws_out.cell(row=tot_row, column=c, value=None)
        cell.fill = fill_tot
        cell.border = thin_b

    lbl = ws_out.cell(row=tot_row, column=1, value='TOTAL')
    lbl.font = font_tot
    lbl.alignment = al_c

    s_cell = ws_out.cell(row=tot_row, column=17, value=f"=SUM(Q2:Q{tot_row-1})")
    s_cell.font = font_tot_q
    s_cell.alignment = al_r
    s_cell.number_format = '"S/ "#,##0.00'

    ws_out.row_dimensions[tot_row].height = 20
    ws_out.freeze_panes = 'A2'
    ws_out.auto_filter.ref = f"A1:U{tot_row-1}"
    return True

def main():
    base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    print("=" * 60)
    print("  GENERADOR MASIVO DE REPORTES SAP (LA CALERA)")
    print("  Regla: Factura -> Importe Total | RH -> Importe Neto")
    print("=" * 60)
    print(f"Directorio de trabajo: {base_dir}")

    # Búsqueda recursiva de carpetas que contengan archivos de consolidado
    subdirs = []
    for root, dirs, files in os.walk(base_dir):
        if any(p in root for p in ['.git', '.agents', '__pycache__', 'scratch']):
            continue
        if any(f.lower().endswith(('.xlsx', '.xls')) and 'CONSOLIDADO' in f.upper() and not f.startswith('~$') for f in files):
            subdirs.append(root)

    subdirs.sort()

    wb_out = Workbook()
    wb_out.remove(wb_out.active)  # remover hoja vacía por defecto

    procesados = 0
    for sdir in subdirs:
        ok = procesar_departamento(sdir, wb_out)
        if ok:
            procesados += 1

    if procesados == 0:
        print("\n⚠️ No se encontraron carpetas con archivos de consolidado para procesar.")
        return

    out_file = os.path.join(base_dir, 'Masivo Reporte.xlsx')
    try:
        wb_out.save(out_file)
        print("\n" + "=" * 60)
        print(f"🎉 Proceso finalizado exitosamente.")
        print(f"📊 Departamentos procesados: {procesados}")
        print(f"📁 Archivo generado: {out_file}")
        print("=" * 60)
    except PermissionError:
        alt_file = os.path.join(base_dir, 'Masivo Reporte (Actualizado).xlsx')
        wb_out.save(alt_file)
        print("\n" + "=" * 60)
        print(f"⚠️ El archivo 'Masivo Reporte.xlsx' está abierto en Excel y bloqueado.")
        print(f"📁 Se generó una copia con los datos actualizados en: {alt_file}")
        print("=" * 60)

if __name__ == '__main__':
    main()
