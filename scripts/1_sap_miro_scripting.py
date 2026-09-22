#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
SOLUCIÓN 1: SAP GUI SCRIPTING (PYTHON) - RECOMENDADA
======================================================
Automatiza el registro masivo en SAP MIRO leyendo el archivo 'Masivo Reporte.xlsx'.
Conexión directa a la sesión activa de SAP GUI for Windows mediante COM.
Extrae el N° de documento MIRO creado y lo registra en el Excel.

Requisitos:
- Tener SAP GUI abierto con la sesión iniciada.
- Scripting habilitado en SAP Logon (Opciones -> Scripting).
- pip install openpyxl pywin32
"""

import os
import sys
import time
import openpyxl

# Compatibilidad UTF-8 en consola de Windows
if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')
if hasattr(sys.stderr, 'reconfigure'):
    sys.stderr.reconfigure(encoding='utf-8')

try:
    import win32com.client
except ImportError:
    print("❌ Error: Requiere la librería 'pywin32'. Instálala ejecutando:")
    print("   pip install pywin32 openpyxl")
    sys.exit(1)


def conectar_sap():
    """Se conecta a la sesión activa de SAP GUI."""
    print("🔗 Conectando con SAP GUI...")
    try:
        sap_gui_auto = win32com.client.GetObject("SAPGUI")
        application = sap_gui_auto.GetScriptingEngine
        if application.Children.Count == 0:
            raise Exception("No hay conexiones abiertas en SAP GUI. Inicia sesión primero.")
        connection = application.Children(0)
        if connection.Children.Count == 0:
            raise Exception("No hay sesiones activas en la conexión de SAP.")
        session = connection.Children(0)
        print(f"✅ Conectado a SAP exitosamente. Usuario activo: {session.Info.User} en mandante {session.Info.Mandant}")
        return session
    except Exception as e:
        print(f"❌ Error al conectar con SAP: {e}")
        print("\nVerifica que:")
        print("1. SAP GUI esté abierto con sesión iniciada.")
        print("2. El scripting esté habilitado en SAP GUI (Ajustar disposición local Alt+F12 -> Opciones -> Scripting).")
        return None


def registrar_miro(session, fecha_doc, referencia, ruc, importe, texto_pos, oc, modo_simulacion=False):
    """
    Ingresa a MIRO, digita los datos de la factura/RxH y contabiliza.
    Retorna: (exito: bool, mensaje: str, num_miro: str)
    """
    try:
        # 1. Ingresar a la transacción MIRO
        session.findById("wnd[0]/tbar[0]/okcd").text = "/nMIRO"
        session.findById("wnd[0]").sendVKey(0)  # Enter
        time.sleep(0.8)

        # 2. Asignar datos de Cabecera
        # Fecha de documento (ej: 26.08.2026)
        session.findById("wnd[0]/usr/subHEADER_AND_ITEMS:SAPLMR1M:6005/tabsTS/tabpTAB_HEADER/ssubHEADER:SAPLMR1M:6010/ctxtINVFO-BLDAT").text = str(fecha_doc)
        
        # Fecha de contabilización (misma fecha o actual)
        session.findById("wnd[0]/usr/subHEADER_AND_ITEMS:SAPLMR1M:6005/tabsTS/tabpTAB_HEADER/ssubHEADER:SAPLMR1M:6010/ctxtINVFO-BUDAT").text = str(fecha_doc)
        
        # Referencia (RxH o Factura formateada)
        session.findById("wnd[0]/usr/subHEADER_AND_ITEMS:SAPLMR1M:6005/tabsTS/tabpTAB_HEADER/ssubHEADER:SAPLMR1M:6010/txtINVFO-XBLNR").text = str(referencia)
        
        # Importe en soles
        importe_str = f"{float(importe):.2f}"
        session.findById("wnd[0]/usr/subHEADER_AND_ITEMS:SAPLMR1M:6005/tabsTS/tabpTAB_HEADER/ssubHEADER:SAPLMR1M:6010/txtINVFO-WRBTR").text = importe_str
        
        # Moneda PEN
        try:
            session.findById("wnd[0]/usr/subHEADER_AND_ITEMS:SAPLMR1M:6005/tabsTS/tabpTAB_HEADER/ssubHEADER:SAPLMR1M:6010/ctxtINVFO-WAERS").text = "PEN"
        except Exception:
            pass

        # Texto de cabecera / posición
        try:
            session.findById("wnd[0]/usr/subHEADER_AND_ITEMS:SAPLMR1M:6005/tabsTS/tabpTAB_HEADER/ssubHEADER:SAPLMR1M:6010/txtINVFO-BKTXT").text = str(texto_pos)[:25]
        except Exception:
            pass

        # 3. Asignar Orden de Compra (OC)
        if oc:
            session.findById("wnd[0]/usr/subHEADER_AND_ITEMS:SAPLMR1M:6005/ssubITEMS:SAPLMR1M:6011/ctxtRM08M-EBELN").text = str(oc).strip()
            session.findById("wnd[0]").sendVKey(0)  # Enter para traer posiciones
            time.sleep(1.0)

        # 4. Manejar posibles ventanas emergentes de aviso (Pop-ups informativos)
        while session.Children.Count > 1:
            popup = session.Children(session.Children.Count - 1)
            popup.sendVKey(0)  # Enter para cerrar aviso
            time.sleep(0.5)

        # 5. Simular o Contabilizar
        if modo_simulacion:
            # Botón Simular (F7)
            session.findById("wnd[0]/tbar[1]/btn[7]").press()
            time.sleep(1.0)
            msg = session.findById("wnd[0]/sbar").text
            return (True, f"Simulado correctamente: {msg}", "SIMULADO")
        else:
            # Botón Contabilizar (Ctrl + S o botón guardar btn[11])
            session.findById("wnd[0]/tbar[0]/btn[11]").press()
            time.sleep(1.5)

            # Leer la barra de estado de SAP
            msg = session.findById("wnd[0]/sbar").text
            msg_type = session.findById("wnd[0]/sbar").messageType

            if msg_type in ('S', 'I') and any(w in msg.lower() for w in ['creado', 'contabilizado', 'documento']):
                # Extraer el número de documento de la respuesta (ej. 5105600123)
                import re
                num_doc = re.search(r'\b(5\d{9})\b', msg)
                num_miro = num_doc.group(1) if num_doc else msg
                return (True, msg, num_miro)
            elif msg_type == 'E':
                return (False, f"Error SAP: {msg}", "")
            else:
                return (True, msg, "REGISTRADO")

    except Exception as e:
        return (False, f"Excepción durante registro: {e}", "")


def procesar_lote_excel(ruta_excel, modo_simulacion=False):
    """Lee el Excel y ejecuta el registro secuencial en SAP."""
    if not os.path.exists(ruta_excel):
        print(f"❌ No se encontró el archivo: {ruta_excel}")
        return

    session = conectar_sap()
    if not session:
        return

    wb = openpyxl.load_workbook(ruta_excel)
    # Procesar primera hoja de datos (ej. Arequipa Semana 34)
    ws = wb.active
    print(f"\n📑 Procesando hoja: '{ws.title}' ({ws.max_row - 2} registros aprox.)")
    print(f"⚙️ Modo: {'SIMULACIÓN (No guarda datos)' if modo_simulacion else 'REAL (Contabiliza en SAP)'}")
    print("=" * 70)

    # Identificar columnas
    header = [ws.cell(1, c).value for c in range(1, 23)]
    c_fecha = 2   # Col B: Fecha Doc
    c_ref   = 4   # Col D: Referencia
    c_ruc   = 6   # Col F: RUC
    c_tpos  = 12  # Col L: Texto Posición
    c_imp   = 17  # Col Q: Importe MD
    c_oc    = 21  # Col U: N° OC

    # Columna para registrar el resultado en el Excel
    col_res = 22  # Col V: Estado MIRO
    col_doc = 23  # Col W: N° Documento SAP
    ws.cell(1, col_res, value="Estado Registro")
    ws.cell(1, col_doc, value="N° Documento MIRO")

    exitosos = 0
    fallidos = 0

    for r in range(2, ws.max_row + 1):
        id_cab = ws.cell(r, 1).value
        if not id_cab or str(id_cab).upper() == 'TOTAL':
            continue

        fecha_doc = ws.cell(r, c_fecha).value
        referencia = ws.cell(r, c_ref).value
        ruc = ws.cell(r, c_ruc).value
        texto_pos = ws.cell(r, c_tpos).value
        importe = ws.cell(r, c_imp).value
        oc = ws.cell(r, c_oc).value

        print(f"\n[Fila {r}] Registrando OC {oc} | Ref: {referencia} | S/ {importe}...")

        exito, msg, num_miro = registrar_miro(session, fecha_doc, referencia, ruc, importe, texto_pos, oc, modo_simulacion)

        if exito:
            print(f"  ✅ ÉXITO: {msg}")
            ws.cell(r, col_res, value="CONTABILIZADO" if not modo_simulacion else "SIMULADO OK")
            ws.cell(r, col_doc, value=num_miro)
            exitosos += 1
        else:
            print(f"  ❌ ERROR: {msg}")
            ws.cell(r, col_res, value=f"ERROR: {msg[:50]}")
            ws.cell(r, col_doc, value="FALLIDO")
            fallidos += 1

        time.sleep(1.0)  # Pausa de seguridad entre registros

    # Guardar Excel con resultados
    ruta_salida = ruta_excel.replace(".xlsx", " (CON MIRO).xlsx")
    wb.save(ruta_salida)
    print("\n" + "=" * 70)
    print("🏁 PROCESAMIENTO FINALIZADO")
    print(f"✅ Exitosos: {exitosos} | ❌ Fallidos: {fallidos}")
    print(f"📁 Reporte de auditoría guardado en:\n   {ruta_salida}")
    print("=" * 70)


if __name__ == '__main__':
    base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    excel_path = os.path.join(base_dir, "Masivo Reporte.xlsx")

    # Si se pasa el argumento --simular no guarda en SAP, solo prueba campos
    simular = "--simular" in sys.argv
    print("=" * 70)
    print("   SOLUCIÓN 1: REGISTRO AUTOMÁTICO EN SAP MIRO (PYTHON SCRIPTING)")
    print("=" * 70)
    procesar_lote_excel(excel_path, modo_simulacion=simular)
