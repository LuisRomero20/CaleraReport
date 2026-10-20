from __future__ import annotations

import shutil
import sys
import tempfile
import unittest
from collections import Counter
from pathlib import Path
from unittest.mock import Mock, patch

import openpyxl


ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = ROOT / "scripts"
if str(SCRIPTS) not in sys.path:
    sys.path.insert(0, str(SCRIPTS))

from control_ordenes import (  # noqa: E402
    MASTER_HEADERS,
    ConsolidationResult,
    build_workbook,
    consolidate,
    display_document,
    format_document_reference,
    normalize_document,
    read_bd_records,
    _pdf_document_total,
)
from procesar_masivo import formatear_referencia, normalizar_texto_pos  # noqa: E402
from renombrar_documentos_miro import (  # noqa: E402
    analyze_renames,
    apply_renames,
    undo_last_rename,
)


EXPEDIENTES = ROOT / "expedientes"


class ControlOrdersIntegrationTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.all_records = consolidate([EXPEDIENTES], mode="expedientes")
        cls.selected = consolidate([EXPEDIENTES], mode="expedientes", week=41)

    def test_all_expedientes_has_24_verified_records_across_4_provinces(self) -> None:
        self.assertEqual(len(self.all_records.records), 24)
        self.assertEqual(
            Counter(row["provincia"] for row in self.all_records.records),
            Counter({"AREQUIPA": 14, "CHICLAYO": 4, "CHIMBOTE": 3, "ICA": 3}),
        )
        self.assertTrue(all(row["_physical_supported"] for row in self.all_records.records))
        self.assertEqual(len(self.all_records.incidents), 0)
        self.assertTrue(all(row["estatus"] == "PENDIENTE MIRO" for row in self.all_records.records))

    def test_week_41_three_sites_has_11_verified_records(self) -> None:
        self.assertEqual(len(self.selected.records), 11)
        self.assertEqual(
            Counter(row["provincia"] for row in self.selected.records),
            Counter({"CHIMBOTE": 3, "AREQUIPA": 4, "CHICLAYO": 4}),
        )
        self.assertTrue(all(row["_physical_supported"] for row in self.selected.records))
        self.assertTrue(all(row["miro"] == "" for row in self.selected.records))
        self.assertEqual(len({tuple(row[field] for field in ("sem", "provincia", "ruc", "documento", "oc", "hes", "motivo", "valorizado")) for row in self.selected.records}), 11)

    def test_real_chiclayo_relation_is_exact(self) -> None:
        row = next(
            item
            for item in self.selected.records
            if item["provincia"] == "CHICLAYO" and item["documento"] == "E001-41"
        )
        self.assertEqual(row["oc"], "430062018")
        self.assertEqual(row["hes"], "1000094356")
        self.assertEqual(row["motivo"], "SERV. SANEO")
        self.assertEqual(row["proveedor"], "MIRES FIESTAS DIEGO AURELIO DE JESUS")
        self.assertEqual(row["valorizado"], 436.99)
        self.assertEqual(row["estatus"], "PENDIENTE MIRO")

    def test_week_filter_selects_requested_period(self) -> None:
        week_34 = consolidate([EXPEDIENTES], mode="expedientes", week=34)
        self.assertEqual(len(week_34.records), 10)
        self.assertEqual({row["provincia"] for row in week_34.records}, {"AREQUIPA"})
        self.assertTrue(all(bool(row["solped"]) for row in week_34.records))
        week_38 = consolidate([EXPEDIENTES], mode="expedientes", week=38)
        self.assertEqual(len(week_38.records), 1)
        self.assertEqual({row["provincia"] for row in week_38.records}, {"ICA"})

    def test_generated_workbook_has_required_structure_and_live_status(self) -> None:
        workbook = build_workbook(self.selected)
        self.assertEqual(workbook.sheetnames, ["BD", "RESUMEN", "INCIDENCIAS"])
        ws = workbook["BD"]
        self.assertEqual([ws.cell(1, c).value for c in range(1, 24)], MASTER_HEADERS)
        self.assertEqual(ws.freeze_panes, "A2")
        self.assertEqual(ws.auto_filter.ref, "A1:W12")
        self.assertEqual(ws["U2"].value, "")
        self.assertIn("PENDIENTE MIRO", ws["W2"].value)
        self.assertEqual(ws["K2"].value, format_document_reference(self.selected.records[0]["documento"], self.selected.records[0]["tipo_doc"]))
        self.assertEqual(ws["Q2"].value, '=IFERROR(IF(VALUE(P2)=2,2,0),0)')
        self.assertEqual(ws["S2"].value, '=IF(R2="G",2,IF(R2="B",0,1))')
        self.assertEqual(ws["H2"].number_format, "@")
        self.assertEqual(ws["L2"].data_type, "n")
        self.assertIsNotNone(ws["B2"].value)

    def test_manual_fields_survive_regeneration_by_document_key(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            previous = Path(tmp) / "maestro_editado.xlsx"
            build_workbook(self.selected).save(previous)
            workbook = openpyxl.load_workbook(previous)
            ws = workbook["BD"]
            ws["U2"] = "4300039999"
            ws["V2"] = "Validado por Contabilidad"
            ws["T2"] = "Nota manual"
            workbook.save(previous)
            regenerated = consolidate(
                [EXPEDIENTES],
                mode="expedientes",
                week=41,
                previous_master=previous,
            )
            target = next(row for row in regenerated.records if row["documento"] == self.selected.records[0]["documento"] and row["provincia"] == self.selected.records[0]["provincia"])
            self.assertEqual(target["miro"], "4300039999")
            self.assertEqual(target["obs_miro"], "Validado por Contabilidad")
            self.assertIn("Nota manual", target["observaciones"])
            self.assertEqual(target["estatus"], "CERRADO")

    def test_old_sap_helpers_keep_their_contract(self) -> None:
        self.assertEqual(formatear_referencia("E001-41", "RXH"), "02-0E001-0000041")
        self.assertEqual(formatear_referencia("E001-622", "FACTURA"), "01-0E001-0000622")
        self.assertEqual(normalizar_texto_pos("SERV. SANEO", 41), "SERVICIO DE SANEO SEMANA 41")


class ReportAmountTests(unittest.TestCase):
    def test_sap_reference_round_trips_to_document_number(self) -> None:
        self.assertEqual(format_document_reference("E001-47", "RXH"), "02-0E001-0000047")
        self.assertEqual(format_document_reference("E001-622", "FACT"), "01-0E001-0000622")
        self.assertEqual(normalize_document("02-0E001-0000047"), normalize_document("E001-47"))
        self.assertEqual(display_document("01-0E001-0000622"), "E001-622")

    @patch("control_ordenes.PdfReader")
    def test_invoice_uses_importe_total_not_valor_venta(self, pdf_reader: Mock) -> None:
        pdf_reader.return_value.pages = [
            Mock(
                extract_text=Mock(
                    return_value="Valor Venta: S/ 720.00 Importe Total: S/ 849.60"
                )
            )
        ]
        self.assertEqual(_pdf_document_total(Path("factura.pdf"), "FACT"), 849.60)

    @patch("control_ordenes.PdfReader")
    def test_receipt_uses_net_received_total(self, pdf_reader: Mock) -> None:
        pdf_reader.return_value.pages = [
            Mock(
                extract_text=Mock(
                    return_value=(
                        "Total Neto Recibido: LA CALERA 13 Octubre 2026 "
                        "110.00 SOLES"
                    )
                )
            )
        ]
        self.assertEqual(_pdf_document_total(Path("rxh.pdf"), "RXH"), 110.00)

    def test_workbook_uses_control_formulas_and_reference_format(self) -> None:
        record = {
            "id": 1,
            "sem": "41",
            "provincia": "CHICLAYO",
            "tipo_doc": "FACT",
            "documento": "E001-622",
            "valorizado": 849.60,
            "lib_solped": "2",
            "lib_os": "G",
        }
        ws = build_workbook(ConsolidationResult("41", [record], [], 0, ["CHICLAYO"]))["BD"]
        self.assertEqual(ws["K2"].value, "01-0E001-0000622")
        self.assertEqual(ws["Q2"].value, '=IFERROR(IF(VALUE(P2)=2,2,0),0)')
        self.assertEqual(ws["S2"].value, '=IF(R2="G",2,IF(R2="B",0,1))')


class RenameToolTests(unittest.TestCase):
    def _make_master(self, path: Path, miro: str) -> None:
        record = {
            "id": 1,
            "fecha": self.source["fecha"],
            "sem": "41",
            "provincia": "CHICLAYO",
            "solicitante": "ALMACEN",
            "motivo": "SERV. SANEO",
            "solped": "",
            "ruc": self.source["ruc"],
            "proveedor": self.source["proveedor"],
            "tipo_doc": "RXH",
            "documento": "E001-41",
            "valorizado": self.source["valorizado"],
            "detalle": self.source["detalle"],
            "oc": "430062018",
            "hes": "1000094356",
            "lib_solped": "2",
            "solped_liberada": 2,
            "lib_os": "G",
            "os_liberada": 2,
            "observaciones": "",
            "miro": miro,
            "obs_miro": "",
            "estatus": "CERRADO",
            "_physical_supported": True,
            "_evidence": {},
            "_blocking": [],
            "_warnings": [],
        }
        result = ConsolidationResult("41", [record], [], 0, ["CHICLAYO"])
        build_workbook(result).save(path)

    def setUp(self) -> None:
        result = consolidate([EXPEDIENTES], mode="expedientes", week=41)
        self.source = next(row for row in result.records if row["provincia"] == "CHICLAYO" and row["documento"] == "E001-41")

    def test_preview_apply_idempotence_and_undo_on_copies(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp) / "expedientes"
            site = root / "1130_CHICLAYO"
            site.mkdir(parents=True)
            original_source = EXPEDIENTES / "1130_CHICLAYO" / "HES 1000094356 - RXH E001-41 - DIEGO MIRES.pdf"
            copied = site / original_source.name
            shutil.copy2(original_source, copied)
            master = Path(tmp) / "maestro.xlsx"
            self._make_master(master, "4300038610")

            preview = analyze_renames(master, root, scope="hes")
            self.assertEqual([item.status for item in preview], ["LISTO"])
            self.assertEqual(Path(preview[0].new_path).name, f"4300038610 - {original_source.name}")
            operation = apply_renames(preview, root, master_path=master, scope="hes")
            self.assertEqual(operation["entries"][0]["status"], "RENOMBRADO")
            renamed = site / f"4300038610 - {original_source.name}"
            self.assertTrue(renamed.exists())
            self.assertFalse(copied.exists())

            second_preview = analyze_renames(master, root, scope="hes")
            self.assertEqual([item.status for item in second_preview], ["YA RENOMBRADO"])
            undone = undo_last_rename(root)
            self.assertEqual(undone["reverted"], 1)
            self.assertEqual(undone["errors"], [])
            self.assertTrue(copied.exists())
            self.assertFalse(renamed.exists())

    def test_invalid_miro_is_never_ready(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp) / "expedientes"
            (root / "1130_CHICLAYO").mkdir(parents=True)
            master = Path(tmp) / "maestro.xlsx"
            self._make_master(master, "123")
            preview = analyze_renames(master, root, scope="hes")
            self.assertEqual([item.status for item in preview], ["MIRO INVÁLIDO"])


if __name__ == "__main__":
    unittest.main()
